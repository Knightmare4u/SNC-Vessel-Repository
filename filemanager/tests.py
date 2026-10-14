import io
import json
import os
import shutil
import tempfile
import zipfile
from unittest import mock

import openpyxl
from django.conf import settings
from django.contrib.auth.models import User
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import InMemoryUploadedFile, SimpleUploadedFile
from django.db import connection
from django.test import Client, SimpleTestCase, TestCase, override_settings
from django.test.utils import CaptureQueriesContext
from django.urls import reverse

from .admin import FolderPermissionForm
from .models import FileActivity, FolderPermission, UserProfile
from .utils import (
    has_permission,
    normalize_path,
    to_url_path,
    validate_name,
)


class ValidateNameTests(SimpleTestCase):
    def test_valid_names_are_kept_as_is(self):
        for name in [
            'VSL SHARE BS2',
            'My Report #3.pdf',
            '100% done.xlsx',
            '2026 報告',
            '.hidden',
            'v1.2',
            'CONSOLE.txt',
        ]:
            with self.subTest(name=name):
                validate_name(name)

    def test_invalid_names(self):
        for name in [
            '',
            '.',
            '..',
            ' leading',
            'trailing ',
            'ends with dot.',
            'a/b',
            'a\\b',
            'a:b',
            'a*b',
            'a?b',
            'a"b',
            'a<b',
            'a>b',
            'a|b',
            'tab\there',
            'CON',
            'nul.txt',
            'com1',
            'LPT9.log',
            'x' * 256,
        ]:
            with self.subTest(name=name):
                with self.assertRaises(ValidationError):
                    validate_name(name)


class PathHelperTests(SimpleTestCase):
    def test_normalize_path(self):
        cases = {
            '': '/',
            '/': '/',
            '///': '/',
            'a/b/': '/a/b',
            '/a//b': '/a/b',
            'a/../b': '/b',
            '../../etc/passwd': '/etc/passwd',
            'a\\..\\b': '/b',
        }
        for raw, expected in cases.items():
            with self.subTest(raw=raw):
                self.assertEqual(normalize_path(raw), expected)

    def test_to_url_path(self):
        cases = {'': '', '/': '', '///': '', '//VSL SHARE BS/': 'VSL SHARE BS'}
        for raw, expected in cases.items():
            with self.subTest(raw=raw):
                self.assertEqual(to_url_path(raw), expected)


class StorageTestCase(TestCase):
    """Runs each test against an empty temporary FILE_STORAGE_ROOT."""

    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.tmp, ignore_errors=True)
        self.storage = os.path.join(self.tmp, 'vessel_files')
        os.makedirs(self.storage)
        override = override_settings(
            FILE_STORAGE_ROOT=self.storage,
            ARCHIVE_ROOT=os.path.join(self.storage, 'ARCHIVE'),
        )
        override.enable()
        self.addCleanup(override.disable)

    def path(self, rel):
        return os.path.join(self.storage, *rel.split('/'))

    def make_dir(self, rel):
        os.makedirs(self.path(rel), exist_ok=True)

    def make_file(self, rel, content=b'hello'):
        self.make_dir(os.path.dirname(rel))
        with open(self.path(rel), 'wb') as f:
            f.write(content)

    @staticmethod
    def make_user(username, perms=(), superuser=False, password_changed=True):
        if superuser:
            user = User.objects.create_superuser(username, f'{username}@x.com', 'pw')
        else:
            user = User.objects.create_user(username, password='pw')
        UserProfile.objects.create(user=user, password_changed=password_changed)
        for folder_path, permission in perms:
            FolderPermission.objects.create(
                user=user, folder_path=folder_path, permission=permission
            )
        return user

    def login(self, user):
        self.client.force_login(user)

    def post_json(self, url, data):
        return self.client.post(url, json.dumps(data), content_type='application/json')

    def read_response(self, response):
        content = b''.join(response.streaming_content)
        response.close()
        return content


class HasPermissionTests(TestCase):
    def make_user(self, *perms):
        user = User.objects.create_user('u', password='pw')
        for folder_path, permission in perms:
            FolderPermission.objects.create(
                user=user, folder_path=folder_path, permission=permission
            )
        return user

    def test_folder_and_subfolders_are_covered(self):
        user = self.make_user(('/VSL SHARE ALL', 'write'))
        self.assertTrue(has_permission(user, '/VSL SHARE ALL', 'write'))
        self.assertTrue(has_permission(user, '/VSL SHARE ALL/2026/report', 'write'))
        self.assertTrue(has_permission(user, 'VSL SHARE ALL', 'write'))

    def test_sibling_with_same_prefix_is_not_covered(self):
        user = self.make_user(('/VSL SHARE ALL', 'write'))
        self.assertFalse(has_permission(user, '/VSL SHARE ALL OMC', 'read'))
        self.assertFalse(has_permission(user, '/VSL SHARE ALL OMC/sub', 'read'))

    def test_dot_dot_cannot_escape_granted_folder(self):
        user = self.make_user(('/A', 'admin'))
        self.assertFalse(has_permission(user, '/A/../B', 'read'))

    def test_levels(self):
        user = self.make_user(('/R', 'read'), ('/W', 'write'), ('/X', 'admin'))
        self.assertEqual(
            [has_permission(user, '/R', lvl) for lvl in ('read', 'write', 'admin')],
            [True, False, False],
        )
        self.assertEqual(
            [has_permission(user, '/W', lvl) for lvl in ('read', 'write', 'admin')],
            [True, True, False],
        )
        self.assertEqual(
            [has_permission(user, '/X', lvl) for lvl in ('read', 'write', 'admin')],
            [True, True, True],
        )

    def test_permissions_are_additive(self):
        user = self.make_user(('/VSL SHARE READ', 'read'), ('/', 'admin'))
        self.assertTrue(has_permission(user, '/VSL SHARE READ', 'admin'))

    def test_no_permission(self):
        user = self.make_user()
        self.assertFalse(has_permission(user, '/', 'read'))


class PathTraversalTests(StorageTestCase):
    # Attack URLs are written out literally on purpose: they mimic raw requests
    # (e.g. "%2e%2e" for "..") that reverse() would never produce.

    def setUp(self):
        super().setUp()
        with open(os.path.join(self.tmp, 'secret.env'), 'wb') as f:
            f.write(b'SECRET_KEY=top-secret')
        self.make_file('A/a.txt')
        self.make_file('B/b.txt')

    def test_download_cannot_escape_storage_root(self):
        self.login(self.make_user('u', [('/', 'read')]))
        response = self.client.get('/download/%2e%2e/secret.env/')
        self.assertEqual(response.status_code, 404)

    def test_preview_cannot_escape_storage_root(self):
        self.login(self.make_user('u', [('/', 'read')]))
        response = self.client.get('/preview/%2e%2e/secret.env/')
        self.assertEqual(response.status_code, 404)

    def test_download_cannot_jump_to_other_folder(self):
        self.login(self.make_user('u', [('/A', 'read')]))
        response = self.client.get('/download/A/%2e%2e/B/b.txt/')
        self.assertEqual(response.status_code, 403)

    def test_search_cannot_jump_to_other_folder(self):
        self.login(self.make_user('u', [('/A', 'read')]))
        response = self.client.get(
            reverse('search_files'), {'q': 'b', 'folder': 'A/../B'}
        )
        self.assertEqual(response.status_code, 403)

    def test_bulk_download_cannot_escape_storage_root(self):
        self.login(self.make_user('u', [('/', 'read')]))
        # '../secret.env' normalizes to '/secret.env' inside the storage root
        response = self.client.post(
            reverse('bulk_download'), {'paths': ['../secret.env']}
        )
        self.assertEqual(response.status_code, 404)

    def test_create_folder_rejects_invalid_names(self):
        self.login(self.make_user('u', [('/A', 'write')]))
        for name in ['', '.', '..', 'a/b', '../B/x', 'a\\b', '..\\B\\x']:
            with self.subTest(name=name):
                response = self.post_json(
                    reverse('create_folder'), {'folder_path': 'A', 'folder_name': name}
                )
                self.assertEqual(response.status_code, 400)
        self.assertFalse(os.path.exists(self.path('B/x')))

    def test_cannot_delete_storage_root(self):
        self.login(self.make_user('su', superuser=True))
        response = self.post_json(reverse('delete_file'), {'path': 'A/..'})
        self.assertEqual(response.status_code, 400)
        self.assertTrue(os.path.isdir(self.path('A')))

    def test_cannot_rename_storage_root(self):
        self.login(self.make_user('su', superuser=True))
        response = self.post_json(
            reverse('rename_file'), {'old_path': 'A/..', 'new_name': 'x'}
        )
        self.assertEqual(response.status_code, 400)


class FileBrowserTests(StorageTestCase):
    def test_missing_folder_returns_404_and_is_not_created(self):
        self.login(self.make_user('su', superuser=True))
        response = self.client.get(reverse('file_browser', args=['does-not-exist']))
        self.assertEqual(response.status_code, 404)
        self.assertFalse(os.path.exists(self.path('does-not-exist')))

    def test_file_path_returns_404(self):
        self.make_file('A/a.txt')
        self.login(self.make_user('su', superuser=True))
        response = self.client.get(reverse('file_browser', args=['A/a.txt']))
        self.assertEqual(response.status_code, 404)

    def test_storage_root_is_created_if_missing(self):
        shutil.rmtree(self.storage)
        self.login(self.make_user('su', superuser=True))
        self.assertEqual(self.client.get(reverse('file_browser')).status_code, 200)
        self.assertTrue(os.path.isdir(self.storage))

    def test_no_permission_redirects_to_dashboard(self):
        self.make_dir('A')
        self.login(self.make_user('u', [('/B', 'read')]))
        response = self.client.get(reverse('file_browser', args=['A']))
        self.assertRedirects(response, reverse('dashboard'))

    def test_paths_have_no_leading_slash(self):
        self.make_file('A/a.txt')
        self.make_dir('A/sub')
        self.login(self.make_user('u', [('/A', 'read')]))
        # '/browser//A/' is the legacy URL format and must keep working
        for url in [reverse('file_browser', args=['A']), '/browser//A/']:
            with self.subTest(url=url):
                response = self.client.get(url)
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.context['current_path'], 'A')
                self.assertEqual(
                    sorted(item['path'] for item in response.context['items']),
                    ['A/a.txt', 'A/sub'],
                )
                self.assertNotContains(response, '/browser//')

    def test_dashboard_links_have_no_double_slash(self):
        self.make_dir('VSL SHARE BS')
        self.login(self.make_user('u', [('/VSL SHARE BS', 'write')]))
        response = self.client.get(reverse('dashboard'))
        folder_url = reverse('file_browser', args=['VSL SHARE BS'])
        self.assertContains(response, f'href="{folder_url}"')
        self.assertNotContains(response, '/browser//')

    def test_unreadable_entries_do_not_hide_the_rest(self):
        self.make_file('A/a.txt')
        self.make_file('A/bad.txt')
        self.make_file('A/z.txt')
        self.make_dir('A/locked')
        self.make_dir('A/sub')
        real_listdir, real_getsize = os.listdir, os.path.getsize

        def listdir(path):
            if os.path.basename(path) == 'locked':
                raise PermissionError(path)
            return real_listdir(path)

        def getsize(path):
            if os.path.basename(path) == 'bad.txt':
                raise OSError(path)
            return real_getsize(path)

        self.login(self.make_user('su', superuser=True))
        with (
            mock.patch('filemanager.views.os.listdir', side_effect=listdir),
            mock.patch('filemanager.views.os.path.getsize', side_effect=getsize),
        ):
            response = self.client.get(reverse('file_browser', args=['A']))

        items = {item['name']: item for item in response.context['items']}
        self.assertEqual(sorted(items), ['a.txt', 'locked', 'sub', 'z.txt'])
        self.assertEqual(items['locked']['size'], '—')
        self.assertEqual(items['sub']['size'], '0 items')
        self.assertEqual(response.context['file_count'], 2)
        self.assertEqual(response.context['folder_count'], 2)

    def test_superuser_dashboard_root_link(self):
        self.login(self.make_user('su', superuser=True))
        response = self.client.get(reverse('dashboard'))
        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, '/browser//')


class SpecialCharacterUrlTests(StorageTestCase):
    """Names with URL-special characters (e.g. files copied straight onto the share)."""

    def setUp(self):
        super().setUp()
        self.make_file('Report #3/Invoice 100%.pdf')
        self.login(self.make_user('u', [('/Report #3', 'write')]))

    def test_browser_page_uses_encoded_urls(self):
        response = self.client.get(reverse('file_browser', args=['Report #3']))
        download_url = reverse('download_file', args=['Report #3/Invoice 100%.pdf'])
        preview_url = reverse('file_preview', args=['Report #3/Invoice 100%.pdf'])
        self.assertIn('%23', download_url)
        self.assertContains(response, f'data-download-url="{download_url}"')
        self.assertContains(response, f'data-preview-url="{preview_url}"')

        response = self.client.get(download_url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self.read_response(response), b'hello')

    def test_search_results_include_encoded_urls(self):
        response = self.client.get(reverse('search_files'), {'q': 'invoice'})
        [result] = response.json()['results']
        self.assertEqual(
            result['download_url'],
            reverse('download_file', args=['Report #3/Invoice 100%.pdf']),
        )
        self.assertEqual(
            result['preview_url'],
            reverse('file_preview', args=['Report #3/Invoice 100%.pdf']),
        )
        response = self.client.get(result['download_url'])
        self.assertEqual(response.status_code, 200)
        response.close()

    def test_rename_returns_encoded_urls(self):
        self.make_dir('Report #3/Sub')
        response = self.post_json(
            reverse('rename_file'), {'old_path': 'Report #3/Sub', 'new_name': 'Sub2'}
        )
        self.assertEqual(
            response.json()['browser_url'],
            reverse('file_browser', args=['Report #3/Sub2']),
        )


class PreviewSandboxTests(StorageTestCase):
    def setUp(self):
        super().setUp()
        self.login(self.make_user('u', [('/A', 'read')]))

    def preview(self, name, content):
        self.make_file(f'A/{name}', content)
        response = self.client.get(reverse('file_preview', args=[f'A/{name}']))
        if response.streaming:
            self.read_response(response)
        return response

    def test_previews_are_sandboxed(self):
        for name, content in [
            ('evil.html', b'<script>alert(1)</script>'),
            ('evil.htm', b'<script>alert(1)</script>'),
            ('evil.svg', b'<svg xmlns="http://www.w3.org/2000/svg"><script/></svg>'),
            ('evil.xml', b'<x/>'),
            ('pic.png', b'data'),
            ('notes.txt', b'data'),
            ('data.csv', b'a,b'),
            ('unknown.zzz', b'data'),
        ]:
            with self.subTest(name=name):
                response = self.preview(name, content)
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response['Content-Security-Policy'], 'sandbox')

    def test_generated_spreadsheet_html_is_sandboxed(self):
        self.make_dir('A')
        wb = openpyxl.Workbook()
        wb.active.append(['x'])
        wb.save(self.path('A/x.xlsx'))
        response = self.client.get(reverse('file_preview', args=['A/x.xlsx']))
        self.assertEqual(response['Content-Security-Policy'], 'sandbox')

    def test_pdf_video_and_audio_are_not_sandboxed(self):
        for name in ['doc.pdf', 'clip.mp4', 'song.mp3']:
            with self.subTest(name=name):
                response = self.preview(name, b'data')
                self.assertEqual(response.status_code, 200)
                self.assertNotIn('Content-Security-Policy', response)


class FolderPermissionFormTests(StorageTestCase):
    def setUp(self):
        super().setUp()
        self.make_dir('A/sub')
        self.make_dir('B')
        self.make_dir('ARCHIVE/A')
        self.user = self.make_user('u')

    def form(self, folder_path, instance=None):
        data = {'user': self.user.pk, 'folder_path': folder_path, 'permission': 'read'}
        return FolderPermissionForm(data, instance=instance)

    def test_choices_are_existing_folders(self):
        widget = FolderPermissionForm().fields['folder_path'].widget
        self.assertEqual(
            [value for value, _ in widget.choices], ['', '/', '/A', '/A/sub', '/B']
        )

    def test_archive_folder_is_rejected(self):
        for folder_path in ['/ARCHIVE', '/ARCHIVE/A']:
            with self.subTest(folder_path=folder_path):
                form = self.form(folder_path)
                self.assertFalse(form.is_valid())
                self.assertIn('folder_path', form.errors)

    def test_missing_folder_is_rejected(self):
        form = self.form('/missing')
        self.assertFalse(form.is_valid())
        self.assertIn('folder_path', form.errors)

    def test_path_is_normalized(self):
        form = self.form('A/sub/')
        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(form.cleaned_data['folder_path'], '/A/sub')

    def test_unchanged_missing_folder_still_saves(self):
        perm = FolderPermission.objects.create(
            user=self.user, folder_path='/gone', permission='read'
        )
        widget = FolderPermissionForm(instance=perm).fields['folder_path'].widget
        self.assertIn('/gone', [value for value, _ in widget.choices])
        form = self.form('/gone', instance=perm)
        self.assertTrue(form.is_valid(), form.errors)


class ErrorPageTests(StorageTestCase):
    """Preview and download open in a browser tab, so errors are HTML pages."""

    def setUp(self):
        super().setUp()
        self.make_file('A/a.txt')
        self.make_file('B/b.txt')
        self.login(self.make_user('u', [('/A', 'read')]))

    def test_no_permission_shows_403_page(self):
        for view in ['file_preview', 'download_file']:
            with self.subTest(view=view):
                response = self.client.get(reverse(view, args=['B/b.txt']))
                self.assertContains(response, 'do not have permission', status_code=403)
                self.assertTrue(response['Content-Type'].startswith('text/html'))

    def test_missing_file_shows_404_page(self):
        for view in ['file_preview', 'download_file']:
            with self.subTest(view=view):
                response = self.client.get(reverse(view, args=['A/missing.txt']))
                self.assertContains(response, 'does not exist', status_code=404)
                self.assertTrue(response['Content-Type'].startswith('text/html'))

    def test_missing_folder_uses_same_404_page(self):
        response = self.client.get(reverse('file_browser', args=['A/missing']))
        self.assertContains(response, 'does not exist', status_code=404)


class ArchiveReadOnlyTests(StorageTestCase):
    def setUp(self):
        super().setUp()
        self.make_file('ARCHIVE/A/old.txt')
        self.make_file('A/a.txt')
        self.login(self.make_user('su', superuser=True))

    def assertReadOnly(self, response):
        self.assertEqual(response.status_code, 403)
        self.assertEqual(response.json(), {'error': 'The archive is read-only'})

    def test_upload_into_archive(self):
        response = self.client.post(
            reverse('upload_file'),
            {'folder_path': 'ARCHIVE/A', 'files': SimpleUploadedFile('x.txt', b'x')},
        )
        self.assertReadOnly(response)
        self.assertFalse(os.path.exists(self.path('ARCHIVE/A/x.txt')))

    def test_create_folder_in_archive(self):
        response = self.post_json(
            reverse('create_folder'), {'folder_path': 'ARCHIVE', 'folder_name': 'new'}
        )
        self.assertReadOnly(response)

    def test_rename_inside_archive(self):
        response = self.post_json(
            reverse('rename_file'), {'old_path': 'ARCHIVE/A/old.txt', 'new_name': 'x.txt'}
        )
        self.assertReadOnly(response)

    def test_rename_archive_itself(self):
        response = self.post_json(
            reverse('rename_file'), {'old_path': 'ARCHIVE', 'new_name': 'OLD'}
        )
        self.assertReadOnly(response)
        self.assertTrue(os.path.isdir(self.path('ARCHIVE')))

    def test_delete_inside_archive(self):
        response = self.post_json(reverse('delete_file'), {'path': 'ARCHIVE/A/old.txt'})
        self.assertReadOnly(response)
        self.assertTrue(os.path.isfile(self.path('ARCHIVE/A/old.txt')))

    def test_delete_archive_itself(self):
        response = self.post_json(reverse('delete_file'), {'path': 'ARCHIVE'})
        self.assertReadOnly(response)
        self.assertTrue(os.path.isdir(self.path('ARCHIVE')))

    def test_archive_can_still_be_browsed_and_downloaded(self):
        response = self.client.get(reverse('file_browser', args=['ARCHIVE/A']))
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.context['in_archive'])
        response = self.client.get(reverse('download_file', args=['ARCHIVE/A/old.txt']))
        self.assertEqual(response.status_code, 200)
        self.read_response(response)


class PasswordChangeRequiredTests(StorageTestCase):
    def test_redirects_until_password_is_changed(self):
        user = self.make_user('u', [('/', 'read')], password_changed=False)
        self.login(user)
        for url in [
            reverse('dashboard'),
            reverse('file_browser'),
            reverse('search_files') + '?q=a',
        ]:
            with self.subTest(url=url):
                self.assertRedirects(
                    self.client.get(url),
                    reverse('change_password'),
                    fetch_redirect_response=False,
                )

        self.assertEqual(self.client.get(reverse('change_password')).status_code, 200)
        response = self.client.post(
            reverse('change_password'),
            {
                'old_password': 'pw',
                'new_password1': 'a-much-longer-password-2026',
                'new_password2': 'a-much-longer-password-2026',
            },
        )
        self.assertRedirects(response, reverse('dashboard'))
        self.assertEqual(self.client.get(reverse('dashboard')).status_code, 200)

    def test_user_without_profile_must_change_password(self):
        user = User.objects.create_superuser('admin', 'a@x.com', 'pw')
        self.login(user)
        self.assertRedirects(
            self.client.get(reverse('dashboard')),
            reverse('change_password'),
            fetch_redirect_response=False,
        )

    def test_logout_is_allowed(self):
        self.login(self.make_user('u', password_changed=False))
        response = self.client.post(reverse('logout'))
        self.assertRedirects(response, reverse('login'), fetch_redirect_response=False)


class LoginTests(StorageTestCase):
    def setUp(self):
        super().setUp()
        self.make_user('u')

    def test_login_redirects_to_dashboard(self):
        response = self.client.post(reverse('login'), {'username': 'u', 'password': 'pw'})
        self.assertRedirects(response, reverse('dashboard'))

    def test_wrong_password_shows_error(self):
        response = self.client.post(
            reverse('login'), {'username': 'u', 'password': 'wrong'}
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'alert-error')
        self.assertNotIn('_auth_user_id', self.client.session)


# Admin templates use {% static %}; avoid needing a collectstatic manifest in tests
@override_settings(
    STORAGES={
        **settings.STORAGES,
        'staticfiles': {
            'BACKEND': 'django.contrib.staticfiles.storage.StaticFilesStorage'
        },
    }
)
class AdminUserListTests(StorageTestCase):
    def test_query_count_does_not_grow_with_users(self):
        self.login(self.make_user('admin', superuser=True))
        url = reverse('admin:auth_user_changelist')

        def count_queries():
            with CaptureQueriesContext(connection) as ctx:
                self.assertEqual(self.client.get(url).status_code, 200)
            return len(ctx.captured_queries)

        self.make_user('u1', [('/A', 'read'), ('/B', 'write')])
        baseline = count_queries()
        for i in range(2, 7):
            self.make_user(f'u{i}', [('/A', 'read'), ('/B', 'write')])
        self.assertEqual(count_queries(), baseline)

    def run_action(self, action, user):
        return self.client.post(
            reverse('admin:auth_user_changelist'),
            {'action': action, '_selected_action': [user.pk]},
            follow=True,
        )

    def test_grant_full_access_upgrades_existing_root_grant(self):
        self.login(self.make_user('admin', superuser=True))
        user = self.make_user('u1', [('/', 'read')])
        response = self.run_action('grant_full_access', user)
        self.assertEqual(
            FolderPermission.objects.get(user=user, folder_path='/').permission, 'admin'
        )
        self.assertContains(response, 'Granted full access to 1 user(s).')

    def test_grant_read_access_never_downgrades(self):
        self.login(self.make_user('admin', superuser=True))
        user = self.make_user('u1', [('/', 'admin')])
        response = self.run_action('grant_read_access', user)
        self.assertEqual(
            FolderPermission.objects.get(user=user, folder_path='/').permission, 'admin'
        )
        self.assertContains(response, 'left unchanged')

    def test_force_password_change_shows_message(self):
        self.login(self.make_user('admin', superuser=True))
        user = self.make_user('u1')
        response = self.run_action('reset_password_required', user)
        self.assertFalse(UserProfile.objects.get(user=user).password_changed)
        self.assertContains(response, '1 user(s) must change their password.')

    def test_folder_access_column(self):
        self.login(self.make_user('admin', superuser=True))
        self.make_user('u1', [('/A', 'read')])
        response = self.client.get(reverse('admin:auth_user_changelist'))
        self.assertContains(response, '/A (read)')


class FileOperationTests(StorageTestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = cls.make_user('u', [('/A', 'write')])

    def setUp(self):
        super().setUp()
        self.make_file('A/a.txt')
        self.login(self.user)

    def upload(self, folder_path, name='new.txt', content=b'data', client=None, **extra):
        return (client or self.client).post(
            reverse('upload_file'),
            {'folder_path': folder_path, 'files': SimpleUploadedFile(name, content)},
            **extra,
        )

    def test_upload(self):
        response = self.upload('A')
        self.assertEqual(response.status_code, 200)
        with open(self.path('A/new.txt'), 'rb') as f:
            self.assertEqual(f.read(), b'data')

    def test_upload_sanitizes_file_name(self):
        for name, expected in [
            ('My Report #3.pdf', 'My_Report_3.pdf'),
            ('bad:name.txt', 'badname.txt'),
        ]:
            with self.subTest(name=name):
                response = self.upload('A', name=name)
                self.assertEqual(response.json()['uploaded_files'][0]['name'], expected)
                self.assertTrue(os.path.isfile(self.path(f'A/{expected}')))

    def test_upload_duplicate_name_gets_suffix(self):
        self.upload('A', name='dup.txt')
        response = self.upload('A', name='dup.txt')
        self.assertEqual(response.json()['uploaded_files'][0]['name'], 'dup_1.txt')

    def test_upload_name_taken_by_folder_gets_suffix(self):
        self.make_dir('A/dup.txt')
        response = self.upload('A', name='dup.txt')
        self.assertEqual(response.json()['uploaded_files'][0]['name'], 'dup_1.txt')

    def test_upload_never_overwrites_file_created_after_name_check(self):
        # Simulate a concurrent upload creating a.txt after the exists() check
        with mock.patch('filemanager.views.os.path.exists', return_value=False):
            response = self.upload('A', name='a.txt', content=b'new')
        self.assertEqual(response.json()['uploaded_files'][0]['name'], 'a_1.txt')
        with open(self.path('A/a.txt'), 'rb') as f:
            self.assertEqual(f.read(), b'hello')

    def test_failed_upload_leaves_no_partial_file(self):
        def failing_chunks(self, chunk_size=None):
            yield b'part'
            raise OSError('disk full')

        with mock.patch.object(InMemoryUploadedFile, 'chunks', failing_chunks):
            response = self.upload('A', name='big.bin')
        self.assertEqual(response.status_code, 500)
        self.assertFalse(os.path.exists(self.path('A/big.bin')))

    def test_upload_invalid_file_name_returns_json_error(self):
        response = self.client.post(
            reverse('upload_file'),
            {
                'folder_path': 'A',
                'files': [
                    SimpleUploadedFile('ok.txt', b'x'),
                    SimpleUploadedFile('???', b'x'),
                ],
            },
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn('error', response.json())
        # The whole batch is rejected before anything is written
        self.assertFalse(os.path.exists(self.path('A/ok.txt')))

    def test_upload_requires_write_permission(self):
        self.make_dir('B')
        FolderPermission.objects.create(user=self.user, folder_path='/B', permission='read')
        self.assertEqual(self.upload('B').status_code, 403)

    def test_upload_to_missing_folder_returns_404(self):
        response = self.upload('A/missing')
        self.assertEqual(response.status_code, 404)
        self.assertFalse(os.path.exists(self.path('A/missing')))

    def test_upload_requires_csrf_token(self):
        client = Client(enforce_csrf_checks=True)
        client.force_login(self.user)
        self.assertEqual(self.upload('A', client=client).status_code, 403)

        client.get(reverse('file_browser', args=['A']))
        token = client.cookies['csrftoken'].value
        response = self.upload('A', client=client, HTTP_X_CSRFTOKEN=token)
        self.assertEqual(response.status_code, 200)

    def test_download(self):
        response = self.client.get(reverse('download_file', args=['A/a.txt']))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self.read_response(response), b'hello')

    def test_xlsx_preview_escapes_html(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = '<b>S<b>'
        ws.append(['<script>alert(1)</script>'])
        wb.save(self.path('A/x.xlsx'))

        response = self.client.get(reverse('file_preview', args=['A/x.xlsx']))
        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, '<script>alert(1)</script>')
        self.assertContains(response, '&lt;script&gt;alert(1)&lt;/script&gt;')
        self.assertContains(response, '&lt;b&gt;S&lt;b&gt;')

    def test_office_preview_is_logged(self):
        wb = openpyxl.Workbook()
        wb.active.append(['x'])
        wb.save(self.path('A/x.xlsx'))
        self.make_file('A/bad.docx', b'not really a docx')
        for name in ['x.xlsx', 'bad.docx']:
            with self.subTest(name=name):
                self.client.get(reverse('file_preview', args=[f'A/{name}']))
                self.assertTrue(
                    FileActivity.objects.filter(
                        activity_type='view', filepath=f'/A/{name}'
                    ).exists()
                )

    def test_broken_office_file_preview_shows_message(self):
        for name in ['bad.docx', 'bad.xlsx', 'bad.xls']:
            with self.subTest(name=name):
                self.make_file(f'A/{name}', b'not really an office file')
                response = self.client.get(reverse('file_preview', args=[f'A/{name}']))
                self.assertEqual(response.status_code, 422)
                self.assertContains(
                    response, 'This file cannot be previewed', status_code=422
                )

    def test_create_folder(self):
        response = self.post_json(
            reverse('create_folder'), {'folder_path': 'A', 'folder_name': 'New Folder'}
        )
        self.assertEqual(response.status_code, 200)
        self.assertTrue(os.path.isdir(self.path('A/New Folder')))

    def test_create_folder_in_deleted_parent_returns_404(self):
        response = self.post_json(
            reverse('create_folder'), {'folder_path': 'A/gone', 'folder_name': 'new'}
        )
        self.assertEqual(response.status_code, 404)
        self.assertFalse(os.path.exists(self.path('A/gone')))

    def test_rename_folder_keeps_spaces(self):
        self.make_dir('A/VSL SHARE BS')
        response = self.post_json(
            reverse('rename_file'),
            {'old_path': 'A/VSL SHARE BS', 'new_name': 'VSL SHARE BS2'},
        )
        self.assertEqual(response.json()['new_path'], 'A/VSL SHARE BS2')
        self.assertTrue(os.path.isdir(self.path('A/VSL SHARE BS2')))

    def test_rename_folder_rejects_invalid_name(self):
        self.make_dir('A/sub')
        response = self.post_json(
            reverse('rename_file'), {'old_path': 'A/sub', 'new_name': 'a:b'}
        )
        self.assertEqual(response.status_code, 400)
        self.assertTrue(os.path.isdir(self.path('A/sub')))

    def test_rename_file_sanitizes_name(self):
        response = self.post_json(
            reverse('rename_file'), {'old_path': 'A/a.txt', 'new_name': 'My Report?.txt'}
        )
        self.assertEqual(response.json()['new_path'], 'A/My_Report.txt')
        self.assertTrue(os.path.isfile(self.path('A/My_Report.txt')))

    def test_rename_file_invalid_name_returns_json_error(self):
        response = self.post_json(
            reverse('rename_file'), {'old_path': 'A/a.txt', 'new_name': '???'}
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn('error', response.json())
        self.assertTrue(os.path.isfile(self.path('A/a.txt')))

    def test_rename_case_only(self):
        self.make_dir('A/sub')
        for old_path, new_name, expected in [
            ('A/a.txt', 'A.txt', 'A.txt'),
            ('A/sub', 'SUB', 'SUB'),
        ]:
            with self.subTest(old_path=old_path):
                response = self.post_json(
                    reverse('rename_file'), {'old_path': old_path, 'new_name': new_name}
                )
                self.assertEqual(response.status_code, 200)
                self.assertIn(expected, os.listdir(self.path('A')))

    def test_rename_long_names_activity_fits_column(self):
        old_name, new_name = 'x' * 150 + '.txt', 'y' * 150 + '.txt'
        self.make_file(f'A/{old_name}')
        response = self.post_json(
            reverse('rename_file'), {'old_path': f'A/{old_name}', 'new_name': new_name}
        )
        self.assertEqual(response.status_code, 200)
        activity = FileActivity.objects.get(activity_type='rename')
        self.assertLessEqual(len(activity.filename), 255)

    def test_rename_onto_other_existing_file_is_rejected(self):
        self.make_file('A/b.txt', b'keep')
        response = self.post_json(
            reverse('rename_file'), {'old_path': 'A/a.txt', 'new_name': 'b.txt'}
        )
        self.assertEqual(response.status_code, 400)
        with open(self.path('A/b.txt'), 'rb') as f:
            self.assertEqual(f.read(), b'keep')

    def test_rename_returns_url_path(self):
        response = self.post_json(
            reverse('rename_file'), {'old_path': 'A/a.txt', 'new_name': 'b.txt'}
        )
        self.assertEqual(
            response.json()['download_url'], reverse('download_file', args=['A/b.txt'])
        )
        self.assertEqual(
            response.json()['preview_url'], reverse('file_preview', args=['A/b.txt'])
        )
        self.assertTrue(os.path.isfile(self.path('A/b.txt')))

    def test_rename_folder_only_updates_its_own_permissions(self):
        self.make_dir('VSL SHARE ALL')
        self.make_dir('VSL SHARE ALL OMC')
        other = self.make_user(
            'other',
            [
                ('/VSL SHARE ALL', 'read'),
                ('/VSL SHARE ALL/sub', 'read'),
                ('/VSL SHARE ALL OMC', 'read'),
            ],
        )
        self.login(self.make_user('su', superuser=True))
        response = self.post_json(
            reverse('rename_file'), {'old_path': 'VSL SHARE ALL', 'new_name': 'VSL_ALL'}
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            sorted(
                FolderPermission.objects.filter(user=other).values_list(
                    'folder_path', flat=True
                )
            ),
            ['/VSL SHARE ALL OMC', '/VSL_ALL', '/VSL_ALL/sub'],
        )

    def test_delete_requires_admin_permission(self):
        response = self.post_json(reverse('delete_file'), {'path': 'A/a.txt'})
        self.assertEqual(response.status_code, 403)
        self.assertTrue(os.path.isfile(self.path('A/a.txt')))

    def test_delete_moves_to_archive(self):
        FolderPermission.objects.filter(user=self.user).update(permission='admin')
        response = self.post_json(reverse('delete_file'), {'path': 'A/a.txt'})
        self.assertEqual(response.status_code, 200)
        self.assertFalse(os.path.exists(self.path('A/a.txt')))
        self.assertTrue(os.path.isfile(self.path('ARCHIVE/A/a.txt')))

    def test_delete_folder_removes_its_permissions(self):
        FolderPermission.objects.filter(user=self.user).update(permission='admin')
        self.make_dir('A/sub/deeper')
        self.make_dir('A/sub2')
        other = self.make_user(
            'other',
            [('/A/sub', 'read'), ('/A/sub/deeper', 'write'), ('/A/sub2', 'read')],
        )
        response = self.post_json(reverse('delete_file'), {'path': 'A/sub'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            list(
                FolderPermission.objects.filter(user=other).values_list(
                    'folder_path', flat=True
                )
            ),
            ['/A/sub2'],
        )

    def test_failed_folder_delete_keeps_permissions(self):
        FolderPermission.objects.filter(user=self.user).update(permission='admin')
        self.make_dir('A/sub')
        other = self.make_user('other', [('/A/sub', 'read')])
        with mock.patch('filemanager.views.shutil.move', side_effect=OSError):
            response = self.post_json(reverse('delete_file'), {'path': 'A/sub'})
        self.assertEqual(response.status_code, 500)
        self.assertTrue(FolderPermission.objects.filter(user=other).exists())

    def test_bulk_download(self):
        response = self.client.post(reverse('bulk_download'), {'paths': ['A/a.txt']})
        names = zipfile.ZipFile(io.BytesIO(self.read_response(response))).namelist()
        self.assertEqual(names, ['a.txt'])

    def test_bulk_download_puts_selected_items_at_top_level(self):
        self.make_file('A/sub/deep/b.txt')
        response = self.client.post(
            reverse('bulk_download'), {'paths': ['A/a.txt', 'A/sub']}
        )
        names = zipfile.ZipFile(io.BytesIO(self.read_response(response))).namelist()
        self.assertEqual(sorted(names), ['a.txt', 'sub/deep/b.txt'])

    def test_bulk_download_renames_same_named_files(self):
        # e.g. search results from different folders
        self.make_file('A/sub/a.txt')
        response = self.client.post(
            reverse('bulk_download'), {'paths': ['A/a.txt', 'A/sub/a.txt']}
        )
        names = zipfile.ZipFile(io.BytesIO(self.read_response(response))).namelist()
        self.assertEqual(names, ['a.txt', 'a_1.txt'])

    def test_bulk_download_any_path_without_permission_rejects_all(self):
        self.make_file('B/b.txt')
        response = self.client.post(
            reverse('bulk_download'), {'paths': ['A/a.txt', 'B/b.txt']}
        )
        self.assertEqual(response.status_code, 403)
        self.assertFalse(FileActivity.objects.filter(activity_type='download').exists())

    def test_bulk_download_any_missing_path_rejects_all(self):
        response = self.client.post(
            reverse('bulk_download'), {'paths': ['A/a.txt', 'A/missing.txt']}
        )
        self.assertEqual(response.status_code, 404)
        self.assertFalse(FileActivity.objects.filter(activity_type='download').exists())

    def test_search_returns_url_paths(self):
        response = self.client.get(reverse('search_files'), {'q': 'a.txt'})
        self.assertEqual([r['path'] for r in response.json()['results']], ['A/a.txt'])
        response = self.client.get(
            reverse('search_files'), {'q': 'a.txt', 'folder': 'A'}
        )
        self.assertEqual([r['path'] for r in response.json()['results']], ['A/a.txt'])
