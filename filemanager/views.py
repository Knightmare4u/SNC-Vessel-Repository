import json
import logging
import mimetypes
import os
import posixpath
import re
from pathlib import PurePosixPath
from urllib.parse import urlsplit

import mammoth
import openpyxl
import xlrd
import zipstream
from django.conf import settings
from django.contrib import messages
from django.contrib.auth import login, update_session_auth_hash
from django.contrib.auth import views as auth_views
from django.contrib.auth.decorators import login_required
from django.contrib.auth.forms import PasswordChangeForm
from django.core.exceptions import (
    PermissionDenied,
    SuspiciousFileOperation,
    ValidationError,
)
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import Count, Q
from django.http import (
    FileResponse,
    Http404,
    HttpResponse,
    JsonResponse,
    StreamingHttpResponse,
)
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.utils.html import escape
from django.utils.http import (
    content_disposition_header,
    url_has_allowed_host_and_scheme,
)
from django.utils.text import get_valid_filename
from django.views.decorators.http import require_POST
from xlrd.xldate import xldate_as_datetime

from .forms import SignerFormSet, SignForm, SigningCreateForm, SigningFieldsForm
from .models import (
    File,
    FileActivity,
    Folder,
    Signer,
    SigningActivity,
    SigningField,
    SigningRequest,
    UploadSession,
    UserProfile,
)
from .services import filesystem, microsoft_auth
from .services import signing as signing_service
from .utils import (
    display_name,
    format_file_size,
    get_storage_path,
    get_user_permissions,
    has_permission,
    is_in_archive,
    log_activity,
    normalize_path,
    to_url_path,
    validate_name,
)

logger = logging.getLogger(__name__)

ARCHIVE_READ_ONLY = 'The archive is read-only'
SIGNING_LOCKED = 'This file, or a file in this folder, has been sent for signing and cannot be deleted'

FILE_TYPE_CATEGORIES = {
    'document': ['.doc', '.docx', '.txt'],
    'pdf': ['.pdf'],
    'spreadsheet': ['.xls', '.xlsx'],
    'presentation': ['.ppt', '.pptx'],
    'image': ['.jpg', '.jpeg', '.png', '.gif', '.svg', '.webp', '.bmp', '.ico'],
    'video': ['.mp4', '.webm', '.avi', '.mov', '.wmv'],
    'audio': ['.mp3', '.wav', '.ogg', '.flac', '.m4a', '.aac'],
    'zip': ['.zip', '.rar', '.7z', '.tar', '.gz'],
}

CATEGORY_ICONS = {
    'document': '📘',
    'pdf': '📕',
    'spreadsheet': '📗',
    'presentation': '📙',
    'image': '🖼️',
    'video': '🎬',
    'audio': '🎵',
    'zip': '📦',
}

FILE_ICONS = {
    **{
        ext: CATEGORY_ICONS[category]
        for category, exts in FILE_TYPE_CATEGORIES.items()
        for ext in exts
    },
    '.txt': '📄',  # Plain text, not a Word document
}

# How many file names the activity log lists for a zip download
DOWNLOAD_LOG_MAX_NAMES = 5

RANGE_HEADER_RE = re.compile(r'bytes=(\d*)-(\d*)')

SIGNING_LIST_TABS = {
    'to_sign': 'Waiting for me',
    'sent': 'Sent by me',
    'notified': 'Notified',
}

MICROSOFT_UNAVAILABLE = (
    'Microsoft sign-in is unavailable right now. Please try again later.'
)
# Session keys holding a Microsoft sign-in until Microsoft sends the user back
MICROSOFT_FLOW_KEY = 'microsoft_auth_flow'
MICROSOFT_NEXT_KEY = 'microsoft_auth_next'


def get_file_icon(extension):
    return FILE_ICONS.get(extension, '📄')


def is_on_disk(file):
    """The row outlives its file on disk if someone removes it outside the site."""
    if file.file.storage.exists(file.file.name):
        return True
    logger.error(
        "File %s (id %s) is missing on disk: %s", file.name, file.pk, file.file.name
    )
    return False


def get_item_urls(path, is_dir):
    """URLs for a file browser item, built with reverse() so names are encoded."""
    url_path = to_url_path(path)
    if is_dir:
        return {'browser_url': reverse('file_browser', args=[url_path])}
    return {
        'download_url': reverse('download_file', args=[url_path]),
        'preview_url': reverse('file_preview', args=[url_path]),
    }


@login_required
def change_password(request):
    if request.method == 'POST':
        form = PasswordChangeForm(request.user, request.POST)
        if form.is_valid():
            user = form.save()
            update_session_auth_hash(request, user)

            # Mark password as changed
            profile, created = UserProfile.objects.get_or_create(user=user)
            profile.password_changed = True
            profile.save()

            messages.success(request, 'Password changed successfully!')
            return redirect('dashboard')
    else:
        form = PasswordChangeForm(request.user)

    return render(request, 'filemanager/change_password.html', {'form': form})


class LoginView(auth_views.LoginView):
    template_name = 'filemanager/login.html'
    redirect_authenticated_user = True

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['microsoft_login_enabled'] = microsoft_auth.is_enabled()
        return context


def microsoft_login(request):
    if not microsoft_auth.is_enabled():
        raise Http404
    redirect_uri = urlsplit(settings.MICROSOFT_REDIRECT_URI)
    if request.get_host() != redirect_uri.netloc:
        # e.g. the IP address: Microsoft sends users back to the redirect URI's
        # site only, so the flow must be kept in that site's session
        return redirect(
            f'{redirect_uri.scheme}://{redirect_uri.netloc}{request.get_full_path()}'
        )

    next_url = request.GET.get('next', '')
    if not url_has_allowed_host_and_scheme(next_url, {request.get_host()}):
        next_url = ''
    try:
        flow = microsoft_auth.start_flow()
    except Exception:
        logger.exception('Could not start Microsoft sign-in')
        messages.error(request, MICROSOFT_UNAVAILABLE)
        return redirect('login')
    request.session[MICROSOFT_FLOW_KEY] = flow
    request.session[MICROSOFT_NEXT_KEY] = next_url
    return redirect(flow['auth_uri'])


def microsoft_callback(request):
    # One use only, so a reloaded or replayed callback is refused
    flow = request.session.pop(MICROSOFT_FLOW_KEY, None)
    next_url = request.session.pop(MICROSOFT_NEXT_KEY, '')
    if flow is None:
        messages.error(request, 'Your Microsoft sign-in expired. Please try again.')
        return redirect('login')
    try:
        user = microsoft_auth.authenticate(flow, request.GET.dict())
    except microsoft_auth.MicrosoftLoginError as e:
        messages.error(request, str(e))
        return redirect('login')
    except Exception:
        logger.exception('Microsoft sign-in failed')
        messages.error(request, MICROSOFT_UNAVAILABLE)
        return redirect('login')

    login(request, user, backend='django.contrib.auth.backends.ModelBackend')
    request.session[microsoft_auth.SESSION_KEY] = True
    logger.info('User %s signed in with Microsoft', user.username)
    return redirect(next_url or settings.LOGIN_REDIRECT_URL)


@login_required
def dashboard(request):
    user_permissions = get_user_permissions(request.user)

    # Get accessible folders
    accessible_folders = []
    for perm in user_permissions:
        folder_path = perm['folder_path']
        if filesystem.resolve_folder(folder_path):
            accessible_folders.append(
                {
                    'name': os.path.basename(folder_path),
                    'path': to_url_path(folder_path),
                    'permission': perm['permission'],
                }
            )

    # Get recent activities
    recent_activities = FileActivity.objects.filter(user=request.user).order_by(
        '-timestamp'
    )[:10]

    context = {
        'accessible_folders': accessible_folders,
        'user_permissions': user_permissions,
        'recent_activities': recent_activities,
    }
    return render(request, 'filemanager/dashboard.html', context)


@login_required
def file_browser(request, folder_path=''):
    folder_path = to_url_path(folder_path)
    if not has_permission(request.user, folder_path, 'read'):
        messages.error(request, 'You do not have permission to access this folder')
        return redirect('dashboard')

    folder = filesystem.resolve_folder(folder_path)
    if folder is None:
        raise Http404('Folder not found')

    in_archive = is_in_archive(get_storage_path(folder_path))

    subfolders = folder.children.filter(deleted_at__isnull=True).annotate(
        folder_count=Count(
            'children', filter=Q(children__deleted_at__isnull=True), distinct=True
        ),
        file_count=Count(
            'files', filter=Q(files__deleted_at__isnull=True), distinct=True
        ),
    )
    files = folder.files.filter(deleted_at__isnull=True)

    # A file sent for signing links to its request instead of starting a new one
    signing_requests = {
        r.file_id: r for r in SigningRequest.objects.filter(file__in=files)
    }
    visible_requests = set()
    if signing_requests:
        visible_requests = set(
            signing_service.visible_requests(request.user)
            .filter(pk__in=[r.pk for r in signing_requests.values()])
            .values_list('pk', flat=True)
        )

    # Get files and folders
    items = []
    total_size = 0
    file_count = 0
    folder_count = 0

    for subfolder in subfolders:
        rel_path = posixpath.join(folder_path, subfolder.name)
        archive = is_in_archive(get_storage_path(rel_path))
        # Only superusers can open the archive, so others don't see it
        if archive and not request.user.is_superuser:
            continue
        items.append(
            {
                'name': subfolder.name,
                'type': 'folder',
                'path': rel_path,
                'size': f'{subfolder.folder_count + subfolder.file_count} items',
                'size_bytes': 0,
                'modified': subfolder.updated_at,
                'created_at': subfolder.created_at,
                'icon': '📁',
                'is_archive': archive,
            }
        )
        folder_count += 1

    for file in files:
        signing_request = signing_requests.get(file.id)
        items.append(
            {
                'name': file.name,
                'type': 'file',
                'path': posixpath.join(folder_path, file.name),
                'size': format_file_size(file.size),
                'size_bytes': file.size,
                'modified': file.updated_at,
                'created_at': file.created_at,
                'extension': file.extension,
                'icon': get_file_icon(file.extension),
                'signing_request': signing_request,
                'signing_visible': (
                    signing_request is not None
                    and signing_request.pk in visible_requests
                ),
            }
        )
        total_size += file.size
        file_count += 1

    # Sort: folders first, then files
    items.sort(key=lambda x: (x['type'] != 'folder', x['name'].lower()))

    # Breadcrumb
    breadcrumbs = []
    if folder_path:
        parts = folder_path.split('/')
        for i, part in enumerate(parts):
            if part:  # Skip empty parts
                path = '/'.join(parts[: i + 1])
                breadcrumbs.append({'name': part, 'path': path})

    context = {
        'current_path': folder_path,
        'items': items,
        'breadcrumbs': breadcrumbs,
        'can_upload': has_permission(request.user, folder_path, 'write'),
        'can_delete': has_permission(request.user, folder_path, 'admin'),
        'can_rename': has_permission(request.user, folder_path, 'write'),
        'can_create_folder': has_permission(request.user, folder_path, 'write'),
        'in_archive': in_archive,
        'total_size': format_file_size(total_size),
        'file_count': file_count,
        'folder_count': folder_count,
    }
    return render(request, 'filemanager/file_browser.html', context)


@login_required
def download_file(request, file_path):
    file_path = normalize_path(file_path)
    if not has_permission(request.user, os.path.dirname(file_path), 'read'):
        raise PermissionDenied

    file = filesystem.resolve_file(file_path)
    if file is None or not is_on_disk(file):
        raise Http404('File not found')

    # Log download activity
    log_activity(
        request.user,
        file.name,
        file_path,
        'download',
        request.META.get('REMOTE_ADDR'),
        file.size,
    )

    content_type, encoding = mimetypes.guess_type(file.name)
    response = FileResponse(
        file.file.open('rb'),
        content_type=content_type,
        as_attachment=True,
        filename=file.name,
    )
    return response


@login_required
@require_POST
def upload_file(request):
    folder_path = normalize_path(request.POST.get('folder_path', ''))

    if not has_permission(request.user, folder_path, 'write'):
        logger.warning(
            "Upload refused for %s to %s: no write permission",
            request.user,
            folder_path,
        )
        return JsonResponse({'error': 'Permission denied'}, status=403)

    full_path = get_storage_path(folder_path)

    if is_in_archive(full_path):
        logger.warning(
            "Upload refused for %s to %s: folder is archived", request.user, folder_path
        )
        return JsonResponse({'error': ARCHIVE_READ_ONLY}, status=403)

    folder = filesystem.resolve_folder(folder_path)
    if folder is None:
        logger.warning(
            "Upload refused for %s to %s: folder not found", request.user, folder_path
        )
        return JsonResponse({'error': 'Folder not found'}, status=404)

    files = request.FILES.getlist('files')
    uploaded_files = []
    total_size = 0

    # Validate every name first so a bad name rejects the batch before any write
    filenames = []
    for file in files:
        try:
            filenames.append(get_valid_filename(file.name))
        except SuspiciousFileOperation:
            logger.warning(
                "Upload refused for %s to %s: invalid file name %r",
                request.user,
                folder_path,
                file.name,
            )
            return JsonResponse(
                {'error': f'Invalid file name: {file.name}'}, status=400
            )

    for file, filename in zip(files, filenames):
        try:
            saved = filesystem.upload_file(folder, file, filename, request.user)
            filename, file_size = saved.name, saved.size
            total_size += file_size

            # Log upload activity
            relative_path = posixpath.join(folder_path, filename)
            log_activity(
                request.user,
                filename,
                relative_path,
                'upload',
                request.META.get('REMOTE_ADDR'),
                file_size,
            )

            uploaded_files.append(
                {
                    'name': filename,
                    'size': file_size,
                    'formatted_size': format_file_size(file_size),
                }
            )

            logger.info(
                "Uploaded %s (%s bytes) by %s", relative_path, file_size, request.user
            )

        except OSError:
            logger.exception(
                "Upload error for %s to %s by %s", file.name, folder_path, request.user
            )
            return JsonResponse({'error': f'Error uploading {file.name}'}, status=500)

    return JsonResponse(
        {
            'success': True,
            'uploaded_files': uploaded_files,
            'total_size': total_size,
            'formatted_total_size': format_file_size(total_size),
        }
    )


@login_required
def get_upload_progress(request, session_id):
    try:
        session = UploadSession.objects.get(session_id=session_id, user=request.user)
        return JsonResponse(
            {
                'total_files': session.total_files,
                'completed_files': session.completed_files,
                'total_size': session.total_size,
                'uploaded_size': session.uploaded_size,
                'status': session.status,
                'progress': (
                    (session.completed_files / session.total_files * 100)
                    if session.total_files > 0
                    else 0
                ),
            }
        )
    except UploadSession.DoesNotExist:
        return JsonResponse({'error': 'Session not found'}, status=404)


@login_required
@require_POST
def delete_file(request):
    try:
        data = json.loads(request.body)
    except json.JSONDecodeError:
        return JsonResponse({'error': 'Invalid JSON'}, status=400)

    item_path = normalize_path(data.get('path', ''))

    if item_path == '/':
        return JsonResponse({'error': 'Cannot delete the root folder'}, status=400)

    if not has_permission(request.user, os.path.dirname(item_path), 'admin'):
        return JsonResponse({'error': 'Permission denied'}, status=403)

    full_path = get_storage_path(item_path)

    if is_in_archive(full_path):
        return JsonResponse({'error': ARCHIVE_READ_ONLY}, status=403)

    item = filesystem.resolve(item_path)
    if item is None:
        return JsonResponse({'error': 'File/folder not found'}, status=404)

    if signing_service.is_locked(item):
        return JsonResponse({'error': SIGNING_LOCKED}, status=403)

    with transaction.atomic():
        signing_service.delete_drafts(item)
        if isinstance(item, Folder):
            filesystem.delete_folder(item, request.user)
            file_size = 0
        else:
            filesystem.delete_file(item, request.user)
            file_size = item.size

    log_activity(
        request.user,
        item.name,
        item_path,
        'delete',
        request.META.get('REMOTE_ADDR'),
        file_size,
    )

    return JsonResponse({'success': True})


@login_required
def search_files(request):
    query = request.GET.get('q', '').strip()
    if not query:
        return JsonResponse({'results': []})

    file_type = request.GET.get('type', '')
    folder_filter = request.GET.get('folder', '')

    allowed_exts = FILE_TYPE_CATEGORIES.get(file_type, []) if file_type else []

    if folder_filter:
        folder = normalize_path(folder_filter)

        if not has_permission(request.user, folder, 'read'):
            return JsonResponse({'error': 'Permission denied'}, status=403)

        scopes = [folder]
    else:
        # Every grant includes read access
        scopes = [
            normalize_path(perm['folder_path'])
            for perm in get_user_permissions(request.user)
        ]

    folder_paths = filesystem.get_folder_paths()
    files = File.objects.filter(deleted_at__isnull=True, name__icontains=query)

    results = []
    for file in files.order_by('name'):
        if allowed_exts and file.extension not in allowed_exts:
            continue

        folder_path = folder_paths[file.folder_id]
        if not any(PurePosixPath(folder_path).is_relative_to(s) for s in scopes):
            continue

        # Only superusers can open the archive, so others don't find files in it
        in_archive = is_in_archive(get_storage_path(folder_path))
        if in_archive and not request.user.is_superuser:
            continue

        rel_path = to_url_path(posixpath.join(folder_path, file.name))
        results.append(
            {
                'name': file.name,
                'path': rel_path,
                'folder': folder_path,
                'size': file.size,
                'formatted_size': format_file_size(file.size),
                'extension': file.extension,
                'icon': get_file_icon(file.extension),
                'modified': file.updated_at.timestamp(),
                'created': file.created_at.timestamp(),
                **get_item_urls(rel_path, is_dir=False),
            }
        )

    return JsonResponse({'results': results})


@login_required
@require_POST
def create_folder(request):
    try:
        data = json.loads(request.body)
    except json.JSONDecodeError:
        return JsonResponse({'error': 'Invalid JSON'}, status=400)

    folder_path = normalize_path(data.get('folder_path', ''))
    folder_name = data.get('folder_name', '')

    if not has_permission(request.user, folder_path, 'write'):
        return JsonResponse({'error': 'Permission denied'}, status=403)

    try:
        validate_name(folder_name)
    except ValidationError as e:
        return JsonResponse({'error': e.messages[0]}, status=400)

    full_path = get_storage_path(posixpath.join(folder_path, folder_name))

    if is_in_archive(full_path):
        return JsonResponse({'error': ARCHIVE_READ_ONLY}, status=403)

    # The parent may have been deleted while the user's page was open
    parent = filesystem.resolve_folder(folder_path)
    if parent is None:
        return JsonResponse({'error': 'Folder not found'}, status=404)

    try:
        filesystem.create_folder(parent, folder_name, request.user)
    except ValidationError as e:
        return JsonResponse({'error': e.messages[0]}, status=400)
    return JsonResponse({'success': True})


@login_required
@require_POST
def rename_file(request):
    try:
        data = json.loads(request.body)
    except json.JSONDecodeError:
        return JsonResponse({'error': 'Invalid JSON'}, status=400)

    old_path = normalize_path(data.get('old_path', ''))
    new_name = data.get('new_name', '').strip()

    if old_path == '/':
        return JsonResponse({'error': 'Old path is required'}, status=400)

    if not new_name:
        return JsonResponse({'error': 'New name is required'}, status=400)

    folder = os.path.dirname(old_path)

    if not has_permission(request.user, folder, 'write'):
        return JsonResponse({'error': 'Permission denied'}, status=403)

    full_old_path = get_storage_path(old_path)

    item = filesystem.resolve(old_path)
    is_folder = isinstance(item, Folder)

    # Folders keep the user's name; files are sanitized like uploads
    if is_folder:
        try:
            validate_name(new_name)
        except ValidationError as e:
            return JsonResponse({'error': e.messages[0]}, status=400)
    else:
        try:
            new_name = get_valid_filename(new_name)
        except SuspiciousFileOperation:
            return JsonResponse({'error': 'Invalid file name'}, status=400)

    new_path = posixpath.join(folder, new_name)
    full_new_path = get_storage_path(new_path)

    if is_in_archive(full_old_path) or is_in_archive(full_new_path):
        return JsonResponse({'error': ARCHIVE_READ_ONLY}, status=403)

    if item is None:
        return JsonResponse({'error': 'File/Folder not found'}, status=404)

    if new_name == item.name:
        return JsonResponse(
            {
                'success': True,
                'new_path': to_url_path(new_path),
                **get_item_urls(new_path, is_folder),
            }
        )

    try:
        if is_folder:
            filesystem.rename_folder(item, new_name)
        else:
            filesystem.rename_file(item, new_name)
    except ValidationError as e:
        return JsonResponse({'error': e.messages[0]}, status=400)

    modified_str = timezone.localtime(item.updated_at).strftime('%b %d, %Y %H:%M')
    log_activity(
        request.user,
        f'{os.path.basename(old_path)} → {new_name}',
        new_path,
        'rename',
        request.META.get('REMOTE_ADDR'),
    )
    return JsonResponse(
        {
            'success': True,
            'new_path': to_url_path(new_path),
            'modified': modified_str,
            **get_item_urls(new_path, is_folder),
        }
    )


@login_required
@require_POST
def bulk_download(request):
    file_paths = request.POST.getlist('paths')

    if not file_paths:
        return JsonResponse({'error': 'No files selected'}, status=400)

    file_paths = [normalize_path(p) for p in file_paths]

    # All-or-nothing: never hand out a zip that silently leaves files out
    items = []
    for file_path in file_paths:
        if not has_permission(request.user, os.path.dirname(file_path), 'read'):
            raise PermissionDenied
        item = filesystem.resolve(file_path)
        if item is None:
            raise Http404('File not found')
        items.append(item)

    zs = zipstream.ZipStream()
    total_size = 0
    used_names = set()

    def unique_name(name):
        # Search results can hold same-named files from different folders
        base, ext = os.path.splitext(name)
        candidate, counter = name, 1
        while candidate.lower() in used_names:
            candidate = f'{base}_{counter}{ext}'
            counter += 1
        used_names.add(candidate.lower())
        return candidate

    # Selected items go at the top level of the zip, as they appear on screen
    entries = []
    for item in items:
        top_name = unique_name(item.name)
        if isinstance(item, Folder):
            entries += [
                (f'{top_name}/{rel}', file) for rel, file in filesystem.walk(item)
            ]
        else:
            entries.append((top_name, item))

    # Check before streaming: a missing file would break the zip halfway.
    # A list, not a generator, so every missing file gets logged
    if not all([is_on_disk(file) for _, file in entries]):
        raise Http404('File not found')

    for arcname, file in entries:
        zs.add_path(file.file.path, arcname=arcname, recurse=False)
        total_size += file.size

    current_folder = os.path.dirname(file_paths[0].rstrip('/\\')) or '/'

    names = [
        os.path.basename(p.rstrip('/\\')) for p in file_paths[:DOWNLOAD_LOG_MAX_NAMES]
    ]
    name_summary = ', '.join(names)
    if len(file_paths) > DOWNLOAD_LOG_MAX_NAMES:
        name_summary += f', +{len(file_paths) - DOWNLOAD_LOG_MAX_NAMES} more'

    log_activity(
        request.user,
        name_summary,
        current_folder,
        'download',
        request.META.get('REMOTE_ADDR'),
        total_size,
    )

    # A single folder is named after itself, as when zipping one on Windows
    zip_name = 'download.zip'
    if len(items) == 1 and isinstance(items[0], Folder):
        zip_name = f'{items[0].name}.zip'

    response = StreamingHttpResponse(zs, content_type='application/zip')
    response['Content-Disposition'] = content_disposition_header(True, zip_name)
    return response


def needs_sandbox(content_type):
    """Previewed user content gets a sandbox CSP so it cannot run scripts on this
    site's origin. PDF, video and audio are exempt: they cannot run scripts, and
    browsers' built-in viewers/players stop working in a sandboxed document."""
    mime = content_type.split(';')[0]
    return mime != 'application/pdf' and not mime.startswith(('video/', 'audio/'))


def _iter_file_chunk(f, length, chunk_size=8192):
    remaining = length
    try:
        while remaining > 0:
            data = f.read(min(chunk_size, remaining))
            if not data:
                break
            remaining -= len(data)
            yield data
    finally:
        f.close()


@login_required
def file_preview(request, file_path):
    def get_value(cell, wb):
        if cell.ctype == xlrd.XL_CELL_DATE:
            return xldate_as_datetime(cell.value, wb.datemode).strftime(
                '%Y-%m-%d %H:%M:%S'
            )
        return cell.value

    def render_row(row, i):
        tag = 'th' if i == 0 else 'td'
        return ''.join(
            f'<{tag}>{escape(cell) if cell is not None else ""}</{tag}>' for cell in row
        )

    file_path = normalize_path(file_path)
    if not has_permission(request.user, os.path.dirname(file_path), 'read'):
        raise PermissionDenied

    file = filesystem.resolve_file(file_path)
    if file is None or not is_on_disk(file):
        raise Http404('File not found')
    full_path = file.file.path

    # Log once here so every preview type (docx/xlsx return early) is recorded
    log_activity(
        request.user,
        file.name,
        file_path,
        'view',
        request.META.get('REMOTE_ADDR'),
        file.size,
    )

    ext = file.extension
    if ext == '.docx':
        try:
            with open(full_path, 'rb') as f:
                result = mammoth.convert_to_html(f)
        except Exception:
            # Uploaded files are untrusted and the parsers raise many error types
            logger.exception("Preview failed for %s", full_path)
            return render(request, 'filemanager/preview_unavailable.html', status=422)
        html = (
            '<!DOCTYPE html><html><head><meta charset="utf-8">'
            '<style>body{font-family:sans-serif;max-width:900px;margin:40px auto;padding:0 20px}</style>'
            f'</head><body>{result.value}</body></html>'
        )
        response = HttpResponse(html, content_type='text/html; charset=utf-8')
        response['Content-Security-Policy'] = 'sandbox'
        return response

    if ext in ('.xlsx', '.xls'):
        sheets_html = []
        try:
            if ext == '.xls':
                wb = xlrd.open_workbook(full_path)
                for sheet_name in wb.sheet_names():
                    ws = wb.sheet_by_name(sheet_name)
                    rows_html = []
                    for i in range(ws.nrows):
                        row = [get_value(ws.cell(i, j), wb) for j in range(ws.ncols)]
                        if all(cell in ('', None) for cell in row):
                            continue
                        rows_html.append(f'<tr>{render_row(row, i)}</tr>')

                    sheets_html.append(
                        f'<h3>{escape(sheet_name)}</h3>'
                        f'<div class="table-wrap"><table>{"".join(rows_html)}</table></div>'
                    )
            else:
                # An open file, not the path: openpyxl checks a path's extension,
                # and the stored name may have none or another one (e.g. renamed)
                with open(full_path, 'rb') as f:
                    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
                    # read_only keeps reading the file; close it even if reading fails
                    try:
                        for sheet_name in wb.sheetnames:
                            ws = wb[sheet_name]
                            rows_html = []
                            for i, row in enumerate(ws.iter_rows(values_only=True)):
                                if all(cell in ('', None) for cell in row):
                                    continue
                                rows_html.append(f'<tr>{render_row(row, i)}</tr>')

                            sheets_html.append(
                                f'<h3>{escape(sheet_name)}</h3>'
                                f'<div class="table-wrap"><table>{"".join(rows_html)}</table></div>'
                            )
                    finally:
                        wb.close()
        except Exception:
            # Uploaded files are untrusted and the parsers raise many error types
            logger.exception("Preview failed for %s", full_path)
            return render(request, 'filemanager/preview_unavailable.html', status=422)

        html = (
            '<!DOCTYPE html><html><head><meta charset="utf-8"><style>'
            'body{font-family:sans-serif;padding:20px 40px}'
            'h3{margin:24px 0 8px}'
            '.table-wrap{overflow-x:auto}'
            'table{border-collapse:collapse;min-width:100%}'
            'th,td{border:1px solid #d1d5db;padding:6px 12px;white-space:nowrap}'
            'th{background:#f3f4f6;font-weight:600}'
            'tr:nth-child(even) td{background:#f9fafb}'
            '</style></head><body>'
            f'{"".join(sheets_html)}'
            '</body></html>'
        )
        response = HttpResponse(html, content_type='text/html; charset=utf-8')
        response['Content-Security-Policy'] = 'sandbox'
        return response

    content_type, _ = mimetypes.guess_type(file.name)
    if ext in ('.csv', '.tsv', '.log', '.md'):
        content_type = 'text/plain; charset=utf-8'
    elif not content_type:
        content_type = 'application/octet-stream'

    file_size = file.size
    range_match = RANGE_HEADER_RE.fullmatch(request.META.get('HTTP_RANGE', ''))

    if range_match and (range_match.group(1) or range_match.group(2)):
        start_str, end_str = range_match.groups()
        if start_str:
            start = int(start_str)
            end = int(end_str) if end_str else file_size - 1
        else:
            start = max(file_size - int(end_str), 0)
            end = file_size - 1

        if file_size == 0 or start > end or start >= file_size:
            response = HttpResponse(status=416)
            response['Content-Range'] = f'bytes */{file_size}'
            return response

        end = min(end, file_size - 1)
        length = end - start + 1

        f = open(full_path, 'rb')
        f.seek(start)
        response = StreamingHttpResponse(
            _iter_file_chunk(f, length), status=206, content_type=content_type
        )
        response['Content-Range'] = f'bytes {start}-{end}/{file_size}'
        response['Content-Length'] = str(length)
    else:
        response = FileResponse(open(full_path, 'rb'), content_type=content_type)

    response['Accept-Ranges'] = 'bytes'
    if needs_sandbox(content_type):
        response['Content-Security-Policy'] = 'sandbox'
    disposition = content_disposition_header(False, file.name)
    if disposition:
        response['Content-Disposition'] = disposition

    return response


# Signing


def get_visible_request(user, pk):
    """The request if the user created it, signs it or is notified of it; 404
    otherwise, so other users cannot tell whether it exists."""
    return get_object_or_404(signing_service.visible_requests(user), pk=pk)


@login_required
def signing_list(request):
    user = request.user
    tab_querysets = {
        # A subquery, not filter(signers__...): that join would be shared with
        # the signer counts below, so they would only count this user's row
        'to_sign': SigningRequest.objects.filter(
            pk__in=Signer.objects.filter(
                user=user,
                status=Signer.Status.PENDING,
                signing_request__status=SigningRequest.Status.IN_PROGRESS,
            ).values('signing_request')
        ),
        'sent': SigningRequest.objects.filter(created_by=user),
        'notified': SigningRequest.objects.filter(notify_users=user).exclude(
            status=SigningRequest.Status.DRAFT
        ),
    }
    tabs = dict(SIGNING_LIST_TABS)
    # Only people allowed to start signing requests have sent any
    if not user.has_perm(signing_service.REQUEST_SIGNATURES):
        del tabs['sent']
    tab = request.GET.get('tab')
    if tab not in tabs:
        tab = 'to_sign'

    signing_requests = (
        tab_querysets[tab]
        .select_related('created_by', 'file')
        .annotate(
            signer_count=Count('signers', distinct=True),
            signed_count=Count(
                'signers',
                filter=Q(signers__status=Signer.Status.SIGNED),
                distinct=True,
            ),
        )
        .order_by('-created_at')
    )
    page = Paginator(signing_requests, 10).get_page(request.GET.get('page'))

    return render(
        request,
        'filemanager/signing/list.html',
        {
            'tab': tab,
            'tabs': [
                {
                    'key': key,
                    'label': label,
                    'count': tab_querysets[key].count(),
                }
                for key, label in tabs.items()
            ],
            'page': page,
        },
    )


@login_required
def signing_create(request):
    if request.method == 'POST':
        form = SigningCreateForm(request.POST, user=request.user)
        formset = SignerFormSet(request.POST, prefix='signers')
        if form.is_valid() and formset.is_valid():
            signers = [f.cleaned_data['user'] for f in formset.ordered_forms]
            details = {
                'source_path': form.cleaned_data['source_path'],
                'signers': signers,
                'notify_users': list(form.cleaned_data['notify_users']),
                'actor': request.user,
            }
            manual = (
                form.cleaned_data['placement_mode']
                == SigningRequest.PlacementMode.MANUAL
            )
            try:
                if manual:
                    # Saved as a draft; it is sent once the boxes are placed
                    signing_request = signing_service.create_draft(
                        placement_mode=SigningRequest.PlacementMode.MANUAL, **details
                    )
                else:
                    signing_request = signing_service.create_and_send(**details)
            except ValidationError as e:
                form.add_error(None, e)
            else:
                if manual:
                    return redirect('signing_fields', pk=signing_request.pk)
                return _sent_redirect(request, signing_request)
        source_path = form.data.get('source_path', '')
    else:
        source_path = normalize_path(request.GET.get('path', ''))
        try:
            signing_service.validate_source(request.user, source_path)
        except ValidationError as e:
            messages.error(request, e.messages[0])
            return redirect('file_browser')
        form = SigningCreateForm(
            initial={'source_path': source_path}, user=request.user
        )
        # The sender usually signs first; they can change or remove the row
        formset = SignerFormSet(prefix='signers', initial=[{'user': request.user.pk}])

    folder = to_url_path(posixpath.dirname(source_path))
    return render(
        request,
        'filemanager/signing/create.html',
        {
            'form': form,
            'formset': formset,
            'filename': posixpath.basename(source_path),
            'cancel_url': (
                reverse('file_browser', args=[folder])
                if folder
                else reverse('file_browser')
            ),
        },
    )


def _sent_redirect(request, signing_request):
    # Nothing has been sent to anyone yet when the sender signs first
    if signing_request.signers.first().user == request.user:
        messages.success(request, 'Signing set up. You are the first to sign.')
    else:
        messages.success(request, 'Signing request sent.')
    return redirect('signing_detail', pk=signing_request.pk)


@login_required
def signing_fields(request, pk):
    """Step 2 of manual placement: drag a box onto the PDF for each signer."""
    signing_request = get_object_or_404(
        SigningRequest,
        pk=pk,
        created_by=request.user,
        status=SigningRequest.Status.DRAFT,
    )
    signers = list(signing_request.signers.select_related('user'))

    if request.method == 'POST':
        form = SigningFieldsForm(request.POST)
        if form.is_valid():
            try:
                # Saved even if sending fails, so the boxes are not lost
                signing_service.save_fields(signing_request, form.cleaned_data['boxes'])
                signing_service.send(signing_request, request.user)
            except ValidationError as e:
                form.add_error(None, e)
            else:
                return _sent_redirect(request, signing_request)
    else:
        form = SigningFieldsForm()

    return render(
        request,
        'filemanager/signing/fields.html',
        {
            'signing_request': signing_request,
            'form': form,
            'signers_data': [
                {
                    'id': signer.pk,
                    'order': signer.signing_order,
                    'name': display_name(signer.user),
                }
                for signer in signers
            ],
            'boxes_data': [
                {
                    'signer': field.signer_id,
                    'field_type': field.field_type,
                    'page': field.page,
                    'x': field.x,
                    'y': field.y,
                    'width': field.width,
                    'height': field.height,
                }
                for field in SigningField.objects.filter(signer__in=signers)
            ],
        },
    )


@login_required
@require_POST
def signing_delete(request, pk):
    signing_request = get_object_or_404(
        SigningRequest,
        pk=pk,
        created_by=request.user,
        status=SigningRequest.Status.DRAFT,
    )
    folder = to_url_path(filesystem.get_path(signing_request.file.folder))
    signing_service.delete_draft(signing_request)
    messages.success(request, 'Draft deleted.')
    return redirect(
        reverse('file_browser', args=[folder]) if folder else reverse('file_browser')
    )


@login_required
def signing_sign(request, pk):
    signing_request = get_visible_request(request.user, pk)
    signer = signing_request.signers.filter(user=request.user).first()

    # Participants who cannot sign right now get told why, instead of a 404
    if signing_request.status != SigningRequest.Status.IN_PROGRESS:
        reason = 'This document is no longer waiting for signatures.'
    elif signer is None:
        reason = 'You are not a signer of this document.'
    elif signer.status == Signer.Status.SIGNED:
        reason = 'You have already signed this document.'
    elif signer.status == Signer.Status.WAITING:
        reason = 'It is not your turn to sign yet.'
    else:
        reason = None
    if reason:
        messages.info(request, reason)
        return redirect('signing_detail', pk=signing_request.pk)

    if request.method == 'POST':
        form = SignForm(request.POST, request.FILES, user=request.user)
        if form.is_valid():
            try:
                signing_service.sign(
                    signing_request=signing_request,
                    actor=request.user,
                    signature_png=form.cleaned_data['signature_png'],
                    save_as_default=form.cleaned_data['save_as_default'],
                )
            except ValidationError as e:
                form.add_error(None, e)
            else:
                messages.success(request, 'Your signature has been applied.')
                return redirect('signing_detail', pk=signing_request.pk)
    else:
        form = SignForm(user=request.user)

    return render(
        request,
        'filemanager/signing/sign.html',
        {
            'signing_request': signing_request,
            'signer': signer,
            'signer_count': signing_request.signers.count(),
            'signed_count': signing_request.signers.filter(
                status=Signer.Status.SIGNED
            ).count(),
            'form': form,
            'saved_signature_url': signing_service.saved_signature_data_url(
                request.user
            ),
            'manual': signing_request.placement_mode
            == SigningRequest.PlacementMode.MANUAL,
            # Only this signer's boxes; other people's are not shown
            'my_boxes': list(
                signer.fields.values('field_type', 'page', 'x', 'y', 'width', 'height')
            ),
        },
    )


@login_required
def signing_detail(request, pk):
    signing_request = get_visible_request(request.user, pk)
    signers = list(signing_request.signers.select_related('user'))
    signed_count = sum(s.status == Signer.Status.SIGNED for s in signers)
    completed = signing_request.activities.filter(
        action=SigningActivity.Action.COMPLETED
    ).first()

    return render(
        request,
        'filemanager/signing/detail.html',
        {
            'signing_request': signing_request,
            'signers': signers,
            'signed_count': signed_count,
            'my_turn': any(
                s.user == request.user and s.status == Signer.Status.PENDING
                for s in signers
            ),
            'written_back': completed.detail.get('written_back') if completed else None,
        },
    )


@login_required
def signing_pdf(request, pk):
    signing_request = get_visible_request(request.user, pk)
    # The signed copy holds every signature so far; before anyone signs, the original
    pdf = signing_request.signed_file or signing_request.original_file
    return FileResponse(
        pdf.open('rb'),
        content_type='application/pdf',
        filename=signing_request.filename,
    )
